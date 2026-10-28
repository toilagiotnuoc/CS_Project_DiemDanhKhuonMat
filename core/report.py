# =====================================================================
# FILE: core/report.py
# NHIỆM VỤ: Xuất báo cáo chuyên cần ra file Excel gồm 5 sheet:
#   1. Tong hop hoc vien : mỗi học viên 1 dòng (dòng bị cảnh báo tô ĐỎ)
#   2. Canh bao          : chỉ những học viên bị cảnh báo
#   3. Theo buoi         : mỗi buổi 1 dòng (buổi có tỉ lệ có mặt < 80% tô VÀNG)
#   4. Chi tiet          : mỗi lượt điểm danh 1 dòng (lượt vắng tô ĐỎ)
#   5. Bieu do           : 3 biểu đồ dạng ảnh
# KIẾN THỨC: pandas (to_excel, ExcelWriter, rename), openpyxl (Font, PatternFill,
#            freeze_panes, column_dimensions, iter_rows, add_image), hàm callback
# =====================================================================

import os
from datetime import datetime
from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import charts, config

# Màu nền dùng để tô dòng
NEN_DO = PatternFill("solid", fgColor="FBEBEB")
NEN_VANG = PatternFill("solid", fgColor="F9EEDA")
NEN_XANH_DAM = PatternFill("solid", fgColor="1F4E79")

# Dictionary đổi tên cột (không dấu, dùng trong code) -> tên tiếng Việt có dấu (hiển thị)
TEN_COT = {
    "Ma_HV": "Mã HV", "Ho_ten": "Họ tên", "So_buoi": "Số buổi", "Co_mat": "Có mặt",
    "Vang": "Vắng", "Di_tre": "Đi trễ", "Ve_som": "Về sớm", "Thieu_check_out": "Thiếu check-out",
    "Tong_gio_hoc": "Tổng giờ học", "Ti_le_chuyen_can": "Tỉ lệ chuyên cần", "Canh_bao": "Cảnh báo",
    "Ma_buoi": "Mã buổi", "Ngay": "Ngày", "Chu_de": "Chủ đề", "Si_so": "Sĩ số",
    "Ti_le_co_mat": "Tỉ lệ có mặt", "Check_in": "Check-in", "Check_out": "Check-out",
    "Phut_tre": "Phút trễ", "Phut_ve_som": "Phút về sớm", "Phut_hoc": "Phút học",
    "Ti_le_thoi_luong": "Tỉ lệ thời lượng", "Trang_thai": "Trạng thái",
    "Phuong_thuc_in": "Cách check-in",
}


# ---------------------------------------------------------------------
# HÀM: Trang trí 1 sheet
#   - Dòng tiêu đề: chữ trắng in đậm, nền xanh đậm
#   - Cố định dòng tiêu đề khi cuộn (freeze panes)
#   - Tự chỉnh độ rộng cột theo nội dung dài nhất
#   - Cột tỉ lệ hiển thị dạng phần trăm
# ---------------------------------------------------------------------
def trang_tri_sheet(ws, cot_phan_tram=None):
    if cot_phan_tram is None:
        cot_phan_tram = []

    # B1: Tô dòng tiêu đề (dòng 1)
    for o in ws[1]:
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = NEN_XANH_DAM
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # B2: Cố định dòng 1
    ws.freeze_panes = "A2"

    # B3: Duyệt từng cột: chỉnh độ rộng + định dạng phần trăm nếu cần
    tieu_de = [o.value for o in ws[1]]
    so_thu_tu_cot = 1
    for cot in ws.columns:
        # Tìm độ dài nội dung dài nhất trong cột
        do_dai_lon_nhat = 0
        for o in cot:
            if o.value is not None and len(str(o.value)) > do_dai_lon_nhat:
                do_dai_lon_nhat = len(str(o.value))
        do_rong = min(max(do_dai_lon_nhat + 2, 10), 50)       # rộng từ 10 đến 50
        ws.column_dimensions[get_column_letter(so_thu_tu_cot)].width = do_rong

        # Cột tỉ lệ -> định dạng % (0.875 hiển thị thành 87.5%)
        if tieu_de[so_thu_tu_cot - 1] in cot_phan_tram:
            for o in cot[1:]:                                  # bỏ ô tiêu đề
                o.number_format = "0.0%"
        so_thu_tu_cot = so_thu_tu_cot + 1


# ---------------------------------------------------------------------
# HÀM: Tô màu những dòng thỏa điều kiện
# Input : ws = sheet, ten_cot = cột cần xét, ham_dieu_kien = hàm trả về True/False,
#         mau = màu nền
# ham_dieu_kien là một CALLBACK (hàm được truyền vào hàm khác) như đã học ở buổi Pandas.
# Nhờ vậy 1 hàm to_mau_dong dùng được cho cả 3 sheet với 3 điều kiện khác nhau.
# ---------------------------------------------------------------------
def to_mau_dong(ws, ten_cot, ham_dieu_kien, mau):
    tieu_de = [o.value for o in ws[1]]
    if ten_cot not in tieu_de:
        return
    vi_tri_cot = tieu_de.index(ten_cot)

    for dong in ws.iter_rows(min_row=2):                  # bỏ dòng tiêu đề
        gia_tri = dong[vi_tri_cot].value
        if ham_dieu_kien(gia_tri):
            for o in dong:
                o.fill = mau


# 3 hàm điều kiện dùng làm callback cho to_mau_dong
def co_canh_bao(gia_tri):
    return gia_tri is not None and gia_tri != ""


def ti_le_thap(gia_tri):
    return gia_tri is not None and gia_tri < config.TI_LE_CHUYEN_CAN_TOI_THIEU


def la_vang(gia_tri):
    return gia_tri in config.TRANG_THAI_VANG


# ---------------------------------------------------------------------
# HÀM: Chèn biểu đồ (fig của matplotlib) vào sheet tại ô vi_tri, ví dụ "A3"
# ---------------------------------------------------------------------
def chen_bieu_do(ws, fig, vi_tri):
    bo_nho = BytesIO()                             # file ảo nằm trong bộ nhớ
    fig.savefig(bo_nho, format="png", dpi=110)     # lưu hình vẽ thành ảnh PNG
    plt.close(fig)
    bo_nho.seek(0)                                 # quay về đầu file ảo để đọc
    ws.add_image(Image(bo_nho), vi_tri)


# =====================================================================
# HÀM CHÍNH: Xuất báo cáo
# Input : 3 bảng (tổng hợp học viên, tổng hợp buổi, chi tiết)
#         file: đường dẫn để lưu, hoặc BytesIO() để tải về từ Streamlit,
#               bỏ trống -> tự lưu vào thư mục bao_cao/
# Output: file đã ghi
# =====================================================================
def xuat_bao_cao(tong_hop_hv, tong_hop_buoi, chi_tiet, file=None):
    # B1: Chưa có đường dẫn -> tự đặt tên theo ngày giờ xuất
    if file is None:
        os.makedirs(config.THU_MUC_BAO_CAO, exist_ok=True)
        ten_file = f"bao_cao_chuyen_can_{datetime.now():%Y%m%d_%H%M}.xlsx"
        file = os.path.join(config.THU_MUC_BAO_CAO, ten_file)

    # B2: Chuẩn bị dữ liệu cho từng sheet
    if tong_hop_hv.empty:
        canh_bao = tong_hop_hv
    else:
        canh_bao = tong_hop_hv[tong_hop_hv["Canh_bao"] != ""]

    cot_chi_tiet = ["Ma_buoi", "Ngay", "Chu_de", "Ma_HV", "Ho_ten", "Check_in", "Check_out",
                    "Phut_tre", "Phut_ve_som", "Phut_hoc", "Ti_le_thoi_luong", "Trang_thai",
                    "Phuong_thuc_in"]
    if chi_tiet.empty:
        ct = chi_tiet
    else:
        ct = chi_tiet[cot_chi_tiet].copy()
        ct["Ngay"] = pd.to_datetime(ct["Ngay"]).dt.strftime("%d/%m/%Y")

    buoi = tong_hop_buoi.copy()
    if not buoi.empty:
        buoi["Ngay"] = pd.to_datetime(buoi["Ngay"]).dt.strftime("%d/%m/%Y")

    # B3: Ghi 4 sheet dữ liệu bằng pandas (rename để tên cột có dấu)
    with pd.ExcelWriter(file, engine="openpyxl") as writer:
        tong_hop_hv.rename(columns=TEN_COT).to_excel(writer, sheet_name="Tong hop hoc vien", index=False)
        canh_bao.rename(columns=TEN_COT).to_excel(writer, sheet_name="Canh bao", index=False)
        buoi.rename(columns=TEN_COT).to_excel(writer, sheet_name="Theo buoi", index=False)
        ct.rename(columns=TEN_COT).to_excel(writer, sheet_name="Chi tiet", index=False)

        # B4: Lấy workbook openpyxl ra để trang trí
        wb = writer.book
        trang_tri_sheet(wb["Tong hop hoc vien"], ["Tỉ lệ chuyên cần"])
        trang_tri_sheet(wb["Canh bao"], ["Tỉ lệ chuyên cần"])
        trang_tri_sheet(wb["Theo buoi"], ["Tỉ lệ có mặt"])
        trang_tri_sheet(wb["Chi tiet"], ["Tỉ lệ thời lượng"])

        # B5: Tô màu các dòng cần chú ý
        to_mau_dong(wb["Tong hop hoc vien"], "Cảnh báo", co_canh_bao, NEN_DO)
        to_mau_dong(wb["Theo buoi"], "Tỉ lệ có mặt", ti_le_thap, NEN_VANG)
        to_mau_dong(wb["Chi tiet"], "Trạng thái", la_vang, NEN_DO)

        # B6: Thêm sheet biểu đồ
        ws = wb.create_sheet("Bieu do")
        ws["A1"] = f"Báo cáo chuyên cần — xuất lúc {datetime.now():%d/%m/%Y %H:%M}"
        ws["A1"].font = Font(bold=True, size=14)

        if chi_tiet.empty:
            dem = pd.Series(dtype="int64")
        else:
            da_tinh = chi_tiet[~chi_tiet["Trang_thai"].isin(config.TRANG_THAI_CHUA_TINH)]
            dem = da_tinh["Trang_thai"].value_counts()

        chen_bieu_do(ws, charts.bieu_do_ti_le_theo_buoi(tong_hop_buoi), "A3")
        chen_bieu_do(ws, charts.bieu_do_vang_theo_hoc_vien(tong_hop_hv), "A23")
        chen_bieu_do(ws, charts.bieu_do_trang_thai(dem), "A46")

    return file
